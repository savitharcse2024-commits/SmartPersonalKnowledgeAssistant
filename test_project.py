import importlib.util
import shutil
import sqlite3
import tempfile
import unittest
from contextlib import closing
from io import BytesIO
from pathlib import Path

from docx import Document
from streamlit.testing.v1 import AppTest

import knowledge


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATABASE_SPEC = importlib.util.spec_from_file_location(
    "project_database", PROJECT_ROOT / "database.py"
)
database = importlib.util.module_from_spec(DATABASE_SPEC)
DATABASE_SPEC.loader.exec_module(database)


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)

    def tearDown(self):
        self.temporary_directory.cleanup()

    def connect(self, path):
        return closing(sqlite3.connect(path))

    def test_migrates_existing_notes_and_maintains_fts_index(self):
        path = self.root / "legacy.db"
        with self.connect(path) as connection:
            connection.execute(
                """
                CREATE TABLE notes (
                    id INTEGER PRIMARY KEY,
                    title TEXT NOT NULL,
                    content TEXT NOT NULL,
                    tags TEXT,
                    category TEXT,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            connection.execute(
                "INSERT INTO notes (id, title, content) VALUES (17, 'Climate', 'Carbon reduction plan')"
            )
            connection.commit()

        database.create_database(path)
        with self.connect(path) as connection:
            row = connection.execute(
                "SELECT id, title, content, note_date, source FROM notes"
            ).fetchone()
            self.assertEqual(row[:3], (17, "Climate", "Carbon reduction plan"))
            self.assertTrue(row[3])
            self.assertEqual(row[4], "")
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM notes_fts WHERE notes_fts MATCH 'carbon*'"
                ).fetchone()[0],
                1,
            )
            connection.execute(
                "UPDATE notes SET content = 'Updated ocean research' WHERE id = 17"
            )
            connection.commit()
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM notes_fts WHERE notes_fts MATCH 'ocean*'"
                ).fetchone()[0],
                1,
            )

        database.create_database(path)
        with self.connect(path) as connection:
            self.assertEqual(connection.execute("SELECT COUNT(*) FROM notes").fetchone()[0], 1)

    def test_backup_restore_and_invalid_backup_preserves_destination(self):
        source = self.root / "source.db"
        target = self.root / "target.db"
        database.create_database(source)
        database.create_database(target)
        with self.connect(source) as connection:
            connection.execute(
                "INSERT INTO notes (title, content, source) VALUES ('Original', 'Text', 'approved.txt')"
            )
            connection.commit()
        with self.connect(target) as connection:
            connection.execute(
                "INSERT INTO notes (title, content) VALUES ('Current', 'Keep until valid restore')"
            )
            connection.commit()

        backup = database.create_backup(source)
        self.assertTrue(database.is_valid_backup(backup))
        database.restore_backup(backup, target)
        with self.connect(target) as connection:
            restored = connection.execute(
                "SELECT title, content, source FROM notes"
            ).fetchall()
        self.assertEqual(restored, [("Original", "Text", "approved.txt")])

        invalid_backup = b"not a SQLite backup"
        self.assertFalse(database.is_valid_backup(invalid_backup))
        with self.assertRaises(sqlite3.Error):
            database.restore_backup(invalid_backup, target)
        with self.connect(target) as connection:
            self.assertEqual(
                connection.execute("SELECT title FROM notes").fetchall(),
                [("Original",)],
            )

    def test_restore_older_backup_and_upgrade_its_notes_schema(self):
        legacy = self.root / "legacy_backup.db"
        with self.connect(legacy) as connection:
            connection.execute(
                """
                CREATE TABLE notes (
                    id INTEGER PRIMARY KEY,
                    title TEXT NOT NULL,
                    content TEXT NOT NULL,
                    tags TEXT,
                    category TEXT,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            connection.execute(
                "INSERT INTO notes (id, title, content) VALUES (31, 'Legacy', 'Preserve me')"
            )
            connection.commit()

        target = self.root / "target.db"
        database.create_database(target)
        database.restore_backup(legacy.read_bytes(), target)
        with self.connect(target) as connection:
            self.assertEqual(
                connection.execute(
                    "SELECT id, title, content, source FROM notes"
                ).fetchall(),
                [(31, "Legacy", "Preserve me", "")],
            )
            self.assertEqual(
                connection.execute(
                    "SELECT COUNT(*) FROM notes_fts WHERE notes_fts MATCH 'preserve*'"
                ).fetchone()[0],
                1,
            )


class KnowledgeTests(unittest.TestCase):
    class UploadedFile(BytesIO):
        def __init__(self, name, data):
            super().__init__(data)
            self.name = name

        def getvalue(self):
            return super().getvalue()

    def test_txt_pdf_and_docx_extraction(self):
        read = knowledge.read_uploaded_document
        self.assertIn(
            "TXT body",
            read(self.UploadedFile("note.txt", b"\xef\xbb\xbfTXT body")),
        )

        import pymupdf

        pdf = pymupdf.open()
        pdf.new_page().insert_text((72, 72), "PDF body")
        pdf_bytes = pdf.tobytes()
        pdf.close()
        self.assertIn("PDF body", read(self.UploadedFile("note.pdf", pdf_bytes)))

        word = Document()
        word.add_paragraph("DOCX paragraph")
        table = word.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "DOCX table"
        docx_bytes = BytesIO()
        word.save(docx_bytes)
        extracted = read(self.UploadedFile("note.docx", docx_bytes.getvalue()))
        self.assertIn("DOCX paragraph", extracted)
        self.assertIn("DOCX table", extracted)

    def test_offline_hints_and_evidence_passages(self):
        preview = knowledge.make_summary(
            "First sentence. Second sentence. Third sentence. Fourth sentence."
        )
        self.assertTrue(preview.startswith("First sentence. Second sentence. Third sentence."))
        self.assertIn("Simple preview", preview)
        suggestions = knowledge.suggest_tags(
            "Photosynthesis photosynthesis plants energy sunlight", "plants"
        )
        self.assertIn("photosynthesis", suggestions)
        self.assertNotIn("plants", suggestions)
        passages = knowledge.rank_passages(
            "The weather is mild. Photosynthesis stores energy in plants.",
            ["photosyn", "energy"],
        )
        self.assertEqual(passages[0][0], 2)
        self.assertIn("Photosynthesis", passages[0][1])
        clipped = knowledge.rank_passages(
            "Background details. " + ("unrelated " * 150) + "coral ecosystems.",
            ["coral"],
        )
        self.assertLessEqual(len(clipped[0][1]), 503)


class StreamlitWorkflowTests(unittest.TestCase):
    def test_create_search_and_save_summary_in_isolated_app(self):
        with tempfile.TemporaryDirectory() as directory:
            app_directory = Path(directory)
            for filename in ("app.py", "database.py", "knowledge.py"):
                shutil.copy2(PROJECT_ROOT / filename, app_directory / filename)

            app = AppTest.from_file(str(app_directory / "app.py")).run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])
            self.assertTrue(any("Smart Personal Knowledge" in item.value for item in app.title))
            self.assertTrue(any("keyword retrieval" in item.value for item in app.caption))

            app.text_input[0].set_value("Ocean research")
            app.text_area[0].set_value(
                "Ocean research measures coral health. It records water temperature."
            )
            app.text_input[1].set_value("Marine project")
            app.text_input[2].set_value("ocean, coral")
            app.text_input[3].set_value("approved source.txt")
            next(item for item in app.button if item.label == "Save note").click()
            app.run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])

            database_path = app_directory / "knowledge.db"
            with closing(sqlite3.connect(database_path)) as connection:
                row = connection.execute(
                    "SELECT title, category, tags, source FROM notes"
                ).fetchone()
            self.assertEqual(
                row,
                ("Ocean research", "Marine project", "ocean, coral", "approved source.txt"),
            )

            next(
                item for item in app.selectbox if item.label == "Filter by tag"
            ).select("coral")
            app.run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])
            self.assertTrue(
                any("1 matching" in item.label for item in app.selectbox)
            )
            next(
                item for item in app.selectbox if item.label == "Filter by source"
            ).select("approved source.txt")
            app.run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])

            next(
                item
                for item in app.text_input
                if item.label == "Search title, content, topic, tags, or source"
            ).set_value("coral")
            app.run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])
            self.assertTrue(any("1 matching" in item.label for item in app.selectbox))

            question = next(
                item for item in app.text_input if item.label == "Ask a question or enter keywords"
            )
            question.set_value("coral temperature")
            app.run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])
            self.assertTrue(any("Ocean research measures coral health" in item.value for item in app.markdown))

            next(item for item in app.button if item.label == "Save this summary preview").click()
            app.run(timeout=30)
            with closing(sqlite3.connect(database_path)) as connection:
                summary_count = connection.execute(
                    "SELECT COUNT(*) FROM saved_summaries"
                ).fetchone()[0]
            self.assertEqual(summary_count, 1)

            next(item for item in app.button if item.label == "Delete note").click()
            app.run(timeout=30)
            next(item for item in app.button if item.label == "Confirm delete").click()
            app.run(timeout=30)
            self.assertFalse(app.exception, [item.message for item in app.exception])
            with closing(sqlite3.connect(database_path)) as connection:
                self.assertEqual(
                    connection.execute("SELECT COUNT(*) FROM notes").fetchone()[0],
                    0,
                )
                self.assertEqual(
                    connection.execute(
                        "SELECT COUNT(*) FROM saved_summaries"
                    ).fetchone()[0],
                    0,
                )


if __name__ == "__main__":
    unittest.main()

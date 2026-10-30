import csv
import io
import re
import zipfile
from collections import Counter
from pathlib import Path

import pymupdf
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
DOCUMENT_IMPORT_ERRORS = (
    OSError,
    ValueError,
    RuntimeError,
    zipfile.BadZipFile,
    pymupdf.FileDataError,
    PackageNotFoundError,
)
STOP_WORDS = {
    "about", "after", "again", "also", "and", "are", "because", "been",
    "before", "being", "between", "both", "but", "can", "could", "does",
    "each", "for", "from", "have", "into", "is", "it", "its", "more",
    "most", "not", "of", "on", "or", "our", "should", "such", "than",
    "that", "the", "their", "them", "then", "there", "these", "they",
    "this", "those", "through", "under", "use", "using", "was", "were",
    "what", "when", "where", "which", "while", "who", "will", "with",
    "would", "you", "your", "how", "why", "did",
}


def search_terms(text):
    return list(dict.fromkeys(re.findall(r"\w+", text.casefold(), flags=re.UNICODE)))


def fts_expression(terms):
    return " OR ".join('"' + term.replace('"', '""') + '"*' for term in terms)


def read_uploaded_document(uploaded_file, progress_bar=None):
    extension = Path(uploaded_file.name).suffix.lower()
    file_bytes = uploaded_file.getvalue()
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("Files must be 20 MB or smaller.")

    if extension == ".txt":
        return file_bytes.decode("utf-8-sig", errors="replace")

    if extension == ".pdf":
        with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:
            page_count = len(pdf)
            text = []
            for index, page in enumerate(pdf):
                text.append(page.get_text())
                if progress_bar is not None and page_count:
                    progress_bar.progress((index + 1) / page_count)
            return "\n".join(text)

    if extension == ".docx":
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        paragraphs.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        return "\n".join(paragraphs)

    raise ValueError("This file type is not supported.")


def make_csv(notes):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(
        [
            "ID", "Title", "Content", "Category", "Tags", "Created at",
            "Updated at", "Note date", "Source",
        ]
    )
    writer.writerows(notes)
    return output.getvalue().encode("utf-8-sig")


def make_summary(content, sentence_limit=3):
    text = re.sub(r"\s+", " ", content).strip()
    if not text:
        return "There is no text to preview."
    sentences = [part for part in re.split(r"(?<=[.!?])\s+", text) if part]
    summary = " ".join(sentences[:sentence_limit])
    if len(summary) > 700:
        summary = summary[:697].rsplit(" ", 1)[0] + "..."
    if len(sentences) > sentence_limit:
        summary += " (Simple preview of the first sentences.)"
    return summary


def suggest_tags(content, current_tags=""):
    words = re.findall(r"\b[a-zA-Z][a-zA-Z0-9_-]{2,}\b", content.lower())
    existing = {tag.strip().lower() for tag in (current_tags or "").split(",")}
    counts = Counter(
        word for word in words if word not in STOP_WORDS and word not in existing
    )
    return [word for word, count in counts.most_common(8)]


def rank_passages(content, keywords):
    passages = [
        sentence.strip()
        for sentence in re.split(r"(?<=[.!?])\s+|\n+", content)
        if sentence.strip()
    ]
    ranked = []
    for passage in passages:
        words = set(search_terms(passage))
        score = sum(
            1 for keyword in keywords
            if any(word.startswith(keyword) for word in words)
        )
        if score:
            if len(passage) > 600:
                folded = passage.casefold()
                positions = [
                    folded.find(keyword.casefold())
                    for keyword in keywords
                    if folded.find(keyword.casefold()) >= 0
                ]
                start = max(0, min(positions) - 220)
                end = min(len(passage), start + 500)
                excerpt = passage[start:end].strip()
                if start:
                    excerpt = "..." + excerpt
                if end < len(passage):
                    excerpt += "..."
                passage = excerpt
            ranked.append((score, passage))
    ranked.sort(key=lambda item: (-item[0], len(item[1])))
    return ranked
